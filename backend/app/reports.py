# ----------------------------------------------------------------------
# Enterpriseviz
# Copyright (C) 2025 David C Jantz
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
# ----------------------------------------------------------------------
"""
Multi-sheet dependency reports.

The detail pages each show two or three tables of things that depend on the
item being viewed. This module turns the very same data the page rendered into
one formatted .xlsx workbook: a Report Info cover sheet followed by one
worksheet per dependency type.

The data always arrives as the dict returned by the matching ``utils.*_details()``
call, so the report and the page cannot disagree about what depends on what.
What each report contains is declared in :data:`REPORTS` as ``Sheet`` and
``Column`` specs; :func:`write_sheet` is the single place that knows how a sheet
is formatted.
"""
import re
from collections import namedtuple
from collections.abc import Mapping
from dataclasses import dataclass, field as dataclass_field
from datetime import date, datetime
from typing import Any, Callable, Iterable, Optional
from urllib.parse import quote, urlparse

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .export import sanitize_cell
from .models import Portal

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: How long a rendered metadata report stays downloadable. That page is
#: POST-only and its rows come from a live authenticated crawl of the portal,
#: so a GET download link cannot re-derive them; metadata_view stashes the
#: result under metadata_cache_key() and the report view reads it back.
METADATA_REPORT_TTL_SECONDS = 30 * 60


def metadata_cache_key(user_pk, alias):
    """Scoped to the user: a metadata crawl runs under their credentials."""
    return f"metadata_report:{user_pk}:{alias}"


#: Rows written per sheet before the sheet is truncated. The cost of a report is
#: bounded by what the page already does — the detail templates render every one
#: of these rows into HTML with no server-side paging — so this is a backstop
#: against a pathological item, not a routine limit. Raising it costs memory
#: during workbook.save() and nothing else.
MAX_ROWS = 10_000

#: Excel refuses to store a longer string in a single cell.
MAX_CELL_CHARS = 32_767

#: Excel's hyperlink target limit; a longer one corrupts the sheet.
MAX_HYPERLINK_CHARS = 2_000

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10

DATE_FORMAT = "yyyy-mm-dd hh:mm"
INT_FORMAT = "0"

_ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

#: Excel's built-in "Light Gray, Table Style Medium 8". Applying it as a real
#: table (rather than painting cells) is what gives each sheet the filter
#: dropdowns, banded rows and structured references Excel users expect.
TABLE_STYLE = "TableStyleMedium8"

#: The cover sheet is key/value prose rather than tabular data, so it is painted
#: by hand. Its bars are black to match the header row the table style paints on
#: every other sheet.
INFO_BAR_FILL = PatternFill("solid", start_color="FF000000")
INFO_TITLE_FONT = Font(bold=True, color="FFFFFFFF", size=12)
INFO_BAR_FONT = Font(bold=True, color="FFFFFFFF")
INFO_BAR_ALIGNMENT = Alignment(vertical="center")
INFO_BAR_HEIGHT = 20
INFO_COLUMNS = 3

LINK_FONT = Font(color="FF0563C1", underline="single")
LABEL_FONT = Font(bold=True)
UNDERLINE_BORDER = Border(bottom=Side(style="thin", color="FF000000"))
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
TRUNCATED_TAB_COLOR = "FFD83020"

#: Room for the filter dropdown Excel draws inside each header cell.
HEADER_WIDTH_PADDING = 3

INFO_SHEET_TITLE = "Report Info"


# ----------------------------------------------------------------------
# Context and accessors
# ----------------------------------------------------------------------

@dataclass
class ReportContext:
    """Everything a column accessor may need beyond the row itself."""

    #: Portal primary key -> alias. combine_apps() flattens the portal FK to a
    #: bare id, so resolving an alias needs a lookup rather than an attribute.
    portal_aliases: dict = dataclass_field(default_factory=dict)
    user: Any = None
    now: Optional[datetime] = None
    #: Used on the Report Info sheet when details["item"] is None.
    fallback_label: Optional[str] = None

    @classmethod
    def build(cls, user=None, now=None, fallback_label=None):
        return cls(portal_aliases=dict(Portal.objects.values_list("pk", "alias")),
                   user=user, now=now or timezone.now(), fallback_label=fallback_label)


_MISSING = object()


def field(*names, default=None):
    """
    Read the first present attribute or mapping key.

    Dependency rows reach this module in three shapes: model instances from the
    web map page, the flattened dicts combine_apps() produces for the service
    and layer pages, and a third dict shape from service_layer_details() on the
    Layer ID page. Every accessor has to cope with all of them.
    """

    def read(row, ctx):
        for name in names:
            if isinstance(row, Mapping):
                if name in row:
                    return row[name]
            else:
                value = getattr(row, name, _MISSING)
                if value is not _MISSING:
                    return value
        return default

    return read


def portal_alias(row, ctx):
    """
    The portal alias for a row, whichever shape it arrived in.

    Model rows expose ``portal_instance.alias``; service_layer_details() puts
    the alias straight into ``portal_instance``; combine_apps() supplies only
    ``portal_instance_id``, which is why ReportContext carries a pk -> alias map.
    """
    if isinstance(row, Mapping):
        instance = row.get("portal_instance")
        if instance is not None:
            return getattr(instance, "alias", instance)
        return ctx.portal_aliases.get(row.get("portal_instance_id"))

    instance = getattr(row, "portal_instance", None)
    if instance is not None:
        return getattr(instance, "alias", instance)
    return ctx.portal_aliases.get(getattr(row, "portal_instance_id", None))


def _as_text(value):
    return "" if value is None else str(value)


def owner_name(*names):
    """An owner column: the FK object stringifies to the username, dicts hold it already."""

    def read(row, ctx):
        return _as_text(field(*names)(row, ctx))

    return read


def service_urls(row, ctx):
    """service_url is an ArrayField; show every URL, one per line."""
    urls = row.service_url_as_list() if hasattr(row, "service_url_as_list") else field("service_url")(row, ctx)
    if isinstance(urls, (list, tuple)):
        return "\n".join(str(url) for url in urls if url)
    return _as_text(urls)


def first_service_url(row, ctx):
    urls = row.service_url_as_list() if hasattr(row, "service_url_as_list") else None
    return urls[0] if urls else None


def matched_layers(row, ctx):
    """The layer names that matched the search, from the filtered_layers prefetch."""
    return ", ".join(layer.layer_name for layer in getattr(row, "filtered_layers", []) or []
                     if layer.layer_name)


def from_key(key):
    """Rows come straight from a key of the details dict."""

    def rows(details, ctx):
        return details.get(key) or []

    return rows


MapLayerRow = namedtuple("MapLayerRow", "name url type service_item_id webmap_layer_id")


def webmap_layer_rows(details, ctx):
    """
    Normalise Webmap.webmap_layers into rows.

    The field is a JSONField keyed by layer name. Every value the refresh writes
    is a dict of ``url``, ``type``, ``service_item_id`` and ``webmap_layer_id``,
    but the field defaults to an empty dict and nothing constrains what an older
    refresh left behind, so a positional ``[url, type]`` pair and a bare list of
    names are both accepted.
    """
    raw = details.get("layers") or {}
    if isinstance(raw, Mapping):
        items = raw.items()
    else:
        items = ((entry, None) for entry in raw)

    for name, value in items:
        if isinstance(value, Mapping):
            yield MapLayerRow(name, value.get("url"), value.get("type"),
                              value.get("service_item_id"), value.get("webmap_layer_id"))
        elif isinstance(value, (list, tuple)):
            yield MapLayerRow(name, _at(value, 0), _at(value, 1), None, None)
        else:
            yield MapLayerRow(name, None, None, None, None)


def _at(sequence, index):
    try:
        return sequence[index]
    except (IndexError, TypeError):
        return None


# ----------------------------------------------------------------------
# Sheet specification
# ----------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    """
    One worksheet column.

    ``kind`` decides how the value is written: ``text`` and ``longtext`` are
    strings (``longtext`` wraps), ``int`` and ``datetime`` become real numbers
    and dates so the sheet can sort and filter on them, ``url`` links a cell to
    its own value, and ``link`` shows ``accessor`` as the label with ``link``
    supplying the target.
    """

    header: str
    accessor: Callable[[Any, ReportContext], Any]
    kind: str = "text"
    link: Optional[Callable[[Any, ReportContext], Any]] = None
    width: Optional[int] = None


@dataclass(frozen=True)
class Sheet:
    title: str
    columns: tuple
    rows: Callable[[dict, ReportContext], Iterable[Any]]


SERVICE_COLUMNS = (
    Column("Name", field("service_name")),
    Column("URL", service_urls, kind="longtext", width=50),
    Column("Type", field("service_type")),
    Column("Publish Server", field("service_mxd_server")),
    Column("MXD", field("service_mxd")),
    Column("Owner", owner_name("service_owner")),
    Column("Access", field("service_access")),
    Column("Created", field("service_created"), kind="datetime"),
    Column("Modified", field("service_modified"), kind="datetime"),
    Column("Instance", portal_alias),
)

# The layer page matches services by layer name, so name the layers that matched.
LAYER_SERVICE_COLUMNS = SERVICE_COLUMNS + (
    Column("Matched Layer", matched_layers, kind="longtext"),
)

WEBMAP_COLUMNS = (
    Column("Title", field("webmap_title"), kind="link", link=field("webmap_url")),
    Column("Owner", owner_name("webmap_owner")),
    Column("Access", field("webmap_access")),
    Column("Created", field("webmap_created"), kind="datetime"),
    Column("Modified", field("webmap_modified"), kind="datetime"),
    Column("Views", field("webmap_views"), kind="int"),
    Column("Instance", portal_alias),
    Column("URL", field("webmap_url"), kind="url", width=50),
)

APP_COLUMNS = (
    Column("Title", field("app_title"), kind="link", link=field("app_url")),
    Column("Type", field("app_type")),
    Column("Usage", field("usage_type")),
    Column("Owner", owner_name("app_owner", "app_owner__user_username")),
    Column("Access", field("app_access")),
    Column("Created", field("app_created"), kind="datetime"),
    Column("Modified", field("app_modified"), kind="datetime"),
    Column("Views", field("app_views"), kind="int"),
    Column("Instance", portal_alias),
    Column("URL", field("app_url"), kind="url", width=50),
)

# Registered database layers, on the service and layer pages.
DB_LAYER_COLUMNS = (
    Column("Name", field("layer_name")),
    Column("Server", field("layer_server")),
    Column("Database", field("layer_database")),
    Column("Version", field("layer_version")),
    Column("Used By", field("layer_used_by_count"), kind="int"),
    Column("Instance", portal_alias),
)

# The web map page's own layer list, from the webmap_layers JSONField.
MAP_LAYER_COLUMNS = (
    Column("Name", field("name")),
    Column("Type", field("type")),
    Column("Service Item ID", field("service_item_id")),
    Column("Webmap Layer ID", field("webmap_layer_id")),
    Column("URL", field("url"), kind="url", width=50),
)

# service_layer_details() returns its own dict shape, shared by maps and apps.
LAYERID_COLUMNS = (
    Column("Title", field("title"), kind="link", link=field("url")),
    Column("Type", field("type")),
    Column("Usage", field("usage_type")),
    Column("Owner", field("owner")),
    Column("Access", field("access")),
    Column("Created", field("created"), kind="datetime"),
    Column("Modified", field("modified"), kind="datetime"),
    Column("Views", field("views"), kind="int"),
    Column("Item ID", field("id")),
    Column("Instance", portal_alias),
    Column("URL", field("url"), kind="url", width=50),
)

METADATA_COLUMNS = (
    Column("Name", field("title"), kind="link", link=field("url")),
    Column("Type", field("type")),
    Column("Owner", field("owner")),
    Column("Item ID", field("id")),
    Column("Description", field("description")),
    Column("Snippet", field("snippet")),
    Column("Thumbnail", field("thumbnail")),
    Column("Access Information", field("accessInformation")),
    Column("License", field("licenseInfo")),
    Column("Score", field("scoreCompleteness"), kind="int"),
    Column("URL", field("url"), kind="url", width=50),
)

#: What each report contains. The service report also gets a Layers sheet the
#: page itself never renders — service_details() already returns the queryset
#: and the template only counts it.
REPORTS = {
    "map": (
        Sheet("Services", SERVICE_COLUMNS, from_key("services")),
        Sheet("Layers", MAP_LAYER_COLUMNS, webmap_layer_rows),
        Sheet("Apps", APP_COLUMNS, from_key("apps")),
    ),
    "service": (
        Sheet("Web Maps", WEBMAP_COLUMNS, from_key("maps")),
        Sheet("Layers", DB_LAYER_COLUMNS, from_key("layers")),
        Sheet("Apps", APP_COLUMNS, from_key("apps")),
    ),
    "layer": (
        Sheet("Services", LAYER_SERVICE_COLUMNS, from_key("services")),
        Sheet("Web Maps", WEBMAP_COLUMNS, from_key("maps")),
        Sheet("Apps", APP_COLUMNS, from_key("apps")),
    ),
    "layerid": (
        Sheet("Web Maps", LAYERID_COLUMNS, from_key("maps")),
        Sheet("Apps", LAYERID_COLUMNS, from_key("apps")),
    ),
    "metadata": (
        Sheet("Metadata", METADATA_COLUMNS, from_key("metadata")),
    ),
}


# ----------------------------------------------------------------------
# Cell writing
# ----------------------------------------------------------------------

def safe_sheet_title(title, used):
    """
    A worksheet title Excel will accept: 31 characters, none of []:*?/\\,
    and unique within the workbook. ``used`` is a set of lowercased titles
    already taken, and is updated in place.
    """
    clean = _ILLEGAL_SHEET_CHARS.sub("-", str(title or "Sheet")).strip("'")[:31] or "Sheet"
    candidate, counter = clean, 1
    while candidate.lower() in used:
        suffix = f"-{counter}"
        candidate = clean[:31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


_NON_NAME_CHARS = re.compile(r"[^A-Za-z0-9_]")


def safe_table_name(title, used):
    """
    A table name Excel will accept: letters, digits and underscores only, must
    not start with a digit, and unique across the workbook. ``used`` is a set of
    lowercased names already taken, and is updated in place.
    """
    clean = _NON_NAME_CHARS.sub("_", str(title or "Table")).strip("_") or "Table"
    if clean[0].isdigit():
        clean = f"T_{clean}"
    candidate, counter = clean, 1
    while candidate.lower() in used:
        candidate = f"{clean}_{counter}"
        counter += 1
    used.add(candidate.lower())
    return candidate


def add_table(worksheet, first_row, last_row, column_count, used_names, name_hint):
    """
    Turn a header row plus its data into a real Excel table styled with
    :data:`TABLE_STYLE`.

    The table carries its own autofilter, so callers must not also set
    ``worksheet.auto_filter`` over the same range — Excel treats the overlap as
    a repair-on-open error. A table also needs at least one data row, so a sheet
    that came back empty is left as a plain header row.
    """
    if last_row <= first_row:
        return None

    ref = f"A{first_row}:{get_column_letter(column_count)}{last_row}"
    table = Table(displayName=safe_table_name(name_hint, used_names), ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=TABLE_STYLE, showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    worksheet.add_table(table)
    return table


def safe_text(value):
    """Everything a string has to survive before openpyxl will store it."""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS - 1] + "…"
    return sanitize_cell(text)


def excel_datetime(value):
    """
    openpyxl refuses timezone-aware datetimes. Model fields come back naive
    under USE_TZ = False, but utils.epoch_to_datetime() builds aware ones that
    flow into these fields during a refresh, so normalise here.
    """
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.make_naive(timezone.localtime(value))
        return value
    if isinstance(value, date):
        return value
    return None


def linkable(href):
    """
    Only http(s) targets under Excel's length limit get a hyperlink. These URLs
    come from portal items and are user-controlled; a file:// or javascript:
    target in a workbook someone opens is a real hazard.
    """
    if not href:
        return False
    text = str(href)
    if len(text) > MAX_HYPERLINK_CHARS:
        return False
    try:
        return urlparse(text).scheme in ("http", "https")
    except ValueError:
        return False


def _apply(cell, column, row, value, ctx):
    """Write one value into one cell according to the column's kind."""
    if value is None or value == "":
        return

    if column.kind == "int":
        try:
            cell.value = int(value)
            cell.number_format = INT_FORMAT
            return
        except (TypeError, ValueError):
            cell.value = safe_text(value)
            return

    if column.kind == "datetime":
        moment = excel_datetime(value)
        if moment is None:
            cell.value = safe_text(value)
            return
        cell.value = moment
        cell.number_format = DATE_FORMAT
        return

    if column.kind in ("url", "link"):
        cell.value = safe_text(value)
        href = value if column.kind == "url" else (column.link(row, ctx) if column.link else None)
        if linkable(href):
            cell.hyperlink = str(href)
            cell.font = LINK_FONT
        return

    cell.value = safe_text(value)
    if column.kind == "longtext":
        cell.alignment = WRAP_ALIGNMENT


SheetResult = namedtuple("SheetResult", "title rows truncated")


def _display_width(value):
    if value is None:
        return 0
    if isinstance(value, (datetime, date)):
        return len(DATE_FORMAT)
    longest = max((len(part) for part in str(value).split("\n")), default=0)
    return longest


def write_sheet(workbook, sheet, rows, ctx, used_titles, used_names):
    """
    Write one dependency sheet as a styled Excel table, with a frozen header row
    and content-fitted column widths. Returns a SheetResult for the cover sheet.

    Header cells are left unpainted: the table style supplies the header fill and
    the banded rows, and an explicit fill would sit on top of it.
    """
    worksheet = workbook.create_sheet(safe_sheet_title(sheet.title, used_titles))
    widths = []

    for index, column in enumerate(sheet.columns, start=1):
        worksheet.cell(row=1, column=index, value=safe_text(column.header))
        widths.append(len(column.header) + HEADER_WIDTH_PADDING)

    written = 0
    truncated = False
    for row in rows:
        if written >= MAX_ROWS:
            truncated = True
            break
        written += 1
        for index, column in enumerate(sheet.columns, start=1):
            cell = worksheet.cell(row=written + 1, column=index)
            _apply(cell, column, row, column.accessor(row, ctx), ctx)
            widths[index - 1] = max(widths[index - 1], _display_width(cell.value))

    worksheet.freeze_panes = "A2"
    add_table(worksheet, 1, written + 1, len(sheet.columns), used_names, sheet.title)

    for index, column in enumerate(sheet.columns, start=1):
        width = column.width or max(MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, widths[index - 1] + 2))
        worksheet.column_dimensions[get_column_letter(index)].width = width

    if truncated:
        worksheet.sheet_properties.tabColor = TRUNCATED_TAB_COLOR

    return SheetResult(worksheet.title, written, truncated)


# ----------------------------------------------------------------------
# Report Info sheet
# ----------------------------------------------------------------------

def _map_info(item, details, ctx):
    return [
        ("Web Map", getattr(item, "webmap_title", None)),
        ("Portal", portal_alias(item, ctx) if item else None),
        ("Item ID", getattr(item, "webmap_id", None)),
        ("URL", getattr(item, "webmap_url", None)),
        ("Owner", _as_text(getattr(item, "webmap_owner", None))),
        ("Created", getattr(item, "webmap_created", None)),
        ("Modified", getattr(item, "webmap_modified", None)),
        ("Views", getattr(item, "webmap_views", None)),
    ]


def _service_info(item, details, ctx):
    return [
        ("Service", getattr(item, "service_name", None)),
        ("Portal", portal_alias(item, ctx) if item else None),
        ("URL", first_service_url(item, ctx) if item else None),
        ("Type", getattr(item, "service_type", None)),
        ("Publish Server", getattr(item, "service_mxd_server", None)),
        ("MXD", getattr(item, "service_mxd", None)),
        ("Owner", _as_text(getattr(item, "service_owner", None))),
        ("Created", getattr(item, "service_created", None)),
        ("Modified", getattr(item, "service_modified", None)),
    ]


def _layer_info(item, details, ctx):
    return [
        ("Layer", getattr(item, "layer_name", None) or ctx.fallback_label),
        ("Server", getattr(item, "layer_server", None)),
        ("Database", getattr(item, "layer_database", None)),
        ("Version", getattr(item, "layer_version", None)),
        ("Portal", portal_alias(item, ctx) if item else None),
        ("Used By", getattr(item, "layer_used_by_count", None)),
    ]


def _layerid_info(item, details, ctx):
    return [
        ("Layer URL", details.get("url_searched") or ctx.fallback_label),
        ("Portal", details.get("instance_alias")),
    ]


def _metadata_info(item, details, ctx):
    return [
        ("Report", "Metadata completeness"),
        ("Portal", details.get("instance_alias")),
    ]


INFO_BUILDERS = {
    "map": _map_info,
    "service": _service_info,
    "layer": _layer_info,
    "layerid": _layerid_info,
    "metadata": _metadata_info,
}

REPORT_TITLES = {
    "map": "Web map dependency report",
    "service": "Service dependency report",
    "layer": "Layer dependency report",
    "layerid": "Layer ID usage report",
    "metadata": "Metadata completeness report",
}


def _write_info(worksheet, kind, details, results, ctx):
    """
    The cover sheet: what this report is, what it describes, and what is in it.

    Two labelled sections, Summary and Contents, painted by hand — the Contents
    listing is a short summary of the workbook rather than data anyone would
    filter or sort, so it is deliberately not an Excel table like the dependency
    sheets are.
    """
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 70
    worksheet.column_dimensions["C"].width = 22
    worksheet.sheet_view.showGridLines = False

    def bar(row, text, font):
        """A full-width black band introducing a section."""
        for column in range(1, INFO_COLUMNS + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.fill = INFO_BAR_FILL
            cell.alignment = INFO_BAR_ALIGNMENT
        cell = worksheet.cell(row=row, column=1, value=safe_text(text))
        cell.font = font
        worksheet.row_dimensions[row].height = INFO_BAR_HEIGHT

    def pair(row, label, value):
        worksheet.cell(row=row, column=1, value=safe_text(label)).font = LABEL_FONT
        cell = worksheet.cell(row=row, column=2)
        moment = excel_datetime(value)
        if moment is not None:
            cell.value = moment
            cell.number_format = DATE_FORMAT
        elif isinstance(value, int) and not isinstance(value, bool):
            cell.value = value
            cell.number_format = INT_FORMAT
        elif value not in (None, ""):
            cell.value = safe_text(value)
            if linkable(value):
                cell.hyperlink = str(value)
                cell.font = LINK_FONT

    row = 1
    bar(row, REPORT_TITLES.get(kind, "Dependency report"), INFO_TITLE_FONT)
    row += 2

    bar(row, "Summary", INFO_BAR_FONT)
    row += 1

    item = details.get("item")
    for label, value in INFO_BUILDERS[kind](item, details, ctx):
        pair(row, label, value)
        row += 1
    if item is None and kind not in ("layerid", "metadata"):
        pair(row, "Note", "No item matched exactly; the report covers the requested name.")
        row += 1

    # A blank line separates what the report describes from who produced it.
    row += 1
    pair(row, "Generated", ctx.now)
    row += 1
    pair(row, "Generated by", getattr(ctx.user, "username", None) or "unknown")
    row += 2

    bar(row, "Contents", INFO_BAR_FONT)
    row += 1

    for column, header in enumerate(("Sheet", "Rows", "Truncated"), start=1):
        cell = worksheet.cell(row=row, column=column, value=header)
        cell.font = LABEL_FONT
        cell.border = UNDERLINE_BORDER
    row += 1

    for result in results:
        worksheet.cell(row=row, column=1, value=safe_text(result.title))
        cell = worksheet.cell(row=row, column=2, value=result.rows)
        cell.number_format = INT_FORMAT
        worksheet.cell(row=row, column=3,
                       value=f"Yes - capped at {MAX_ROWS:,}" if result.truncated else "No")
        row += 1



# ----------------------------------------------------------------------
# Public API
# ----------------------------------------------------------------------

def build_dependency_workbook(kind, details, *, user=None, fallback_label=None, now=None):
    """
    Build the workbook for one detail page.

    :param kind: A key of :data:`REPORTS` — "map", "service", "layer",
                 "layerid" or "metadata".
    :param details: Exactly what the matching ``utils.*_details()`` returned,
                    so the report and the page can never disagree.
    :param user: The requesting user, named on the Report Info sheet.
    :param fallback_label: Item name from the URL, used when details["item"] is
                           None — layer_details() returns that when a layer of
                           the name exists but none matches the server,
                           database and version given.
    :return: An openpyxl Workbook.
    """
    if kind not in REPORTS:
        raise ValueError(f"Unknown report kind: {kind}")

    ctx = ReportContext.build(user=user, now=now, fallback_label=fallback_label)

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_titles = set()
    used_names = set()
    info = workbook.create_sheet(safe_sheet_title(INFO_SHEET_TITLE, used_titles))

    results = [write_sheet(workbook, sheet, sheet.rows(details, ctx), ctx,
                           used_titles, used_names)
               for sheet in REPORTS[kind]]

    _write_info(info, kind, details, results, ctx)
    workbook.active = 0
    return workbook


def content_disposition(stem, now=None):
    """
    An attachment header whose filename cannot be broken by a portal-supplied
    title. slugify() drops the quotes, semicolons, newlines and non-ASCII that
    would otherwise escape the header, so both forms end up identical.
    """
    stamp = (now or timezone.now()).strftime("%Y%m%d-%H%M")
    slug = slugify(stem)[:80] or "report"
    name = f"enterpriseviz_{slug}_{stamp}.xlsx"
    return f"attachment; filename=\"{name}\"; filename*=UTF-8''{quote(name)}"
